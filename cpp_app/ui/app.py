from pathlib import Path
import sys
from datetime import date
from decimal import Decimal, InvalidOperation
import tempfile
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtGui import QPainter
from PySide6.QtPrintSupport import QPrinter, QPrinterInfo, QPrintPreviewWidget

from PySide6.QtCore import QDate, Qt
from openpyxl import load_workbook

from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QListWidget,
    QListWidgetItem,
    QTabWidget,
    QComboBox,
    QDateEdit,
    QDialog,
    QFormLayout,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSizePolicy,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from sqlalchemy import or_

from cpp_app.bootstrap import bootstrap_database, refresh_procedure_prices
from cpp_app.config import PRICE_TABLE_PATH
from cpp_app.database import SessionLocal
from cpp_app.models import Procedure
from cpp_app.reports import MONTH_NAMES, export_budget_pdf, export_excel, export_pdf, selected_charged_value
from cpp_app.services import authenticate, create_attendance, create_user, delete_attendance, delete_user, list_users, search_attendances, update_attendance


APP_STYLESHEET = """
QWidget {
    background-color: #f3f5f8;
    color: #1c2530;
    font-family: "Segoe UI";
    font-size: 10pt;
}

QMainWindow, QDialog {
    background-color: #f3f5f8;
}

/* ---------- Abas ---------- */
QTabWidget::pane {
    border: 1px solid #d7dce3;
    border-radius: 8px;
    background-color: #ffffff;
    top: -1px;
}

QTabBar::tab {
    background-color: #e6eaf0;
    color: #4a5568;
    padding: 8px 22px;
    margin-right: 4px;
    border-top-left-radius: 8px;
    border-top-right-radius: 8px;
    font-weight: 600;
}

QTabBar::tab:selected {
    background-color: #1f4e79;
    color: #ffffff;
}

QTabBar::tab:hover:!selected {
    background-color: #d3dbe6;
}

/* ---------- Grupos / cartoes ---------- */
QGroupBox {
    background-color: #ffffff;
    border: 1px solid #d7dce3;
    border-radius: 10px;
    margin-top: 14px;
    padding: 16px 12px 12px 12px;
    font-weight: 600;
    color: #1f4e79;
}

QGroupBox::title {
    subcontrol-origin: margin;
    subcontrol-position: top left;
    left: 12px;
    padding: 0 6px;
    background-color: #ffffff;
}

/* ---------- Campos de entrada ---------- */
QLineEdit, QDateEdit, QComboBox {
    background-color: #ffffff;
    border: 1px solid #cbd3dd;
    border-radius: 6px;
    padding: 5px 8px;
    selection-background-color: #1f4e79;
}

QLineEdit:focus, QDateEdit:focus, QComboBox:focus {
    border: 1px solid #1f4e79;
}

QComboBox::drop-down {
    border: none;
    width: 22px;
}

QListWidget, QTableWidget {
    background-color: #ffffff;
    border: 1px solid #d7dce3;
    border-radius: 6px;
    gridline-color: #e4e8ee;
    selection-background-color: #d7e6f7;
    selection-color: #1c2530;
}

QListWidget::item {
    padding: 4px;
}

QListWidget::item:hover {
    background-color: #eef3fa;
}

QHeaderView::section {
    background-color: #1f4e79;
    color: #ffffff;
    padding: 6px;
    border: none;
    font-weight: 600;
}

QTableWidget {
    alternate-background-color: #f7f9fc;
}

/* ---------- Botoes ---------- */
QPushButton {
    background-color: #ffffff;
    border: 1px solid #c3ccd8;
    border-radius: 6px;
    padding: 7px 16px;
    color: #1c2530;
    font-weight: 600;
}

QPushButton:hover {
    background-color: #eef3fa;
    border: 1px solid #1f4e79;
}

QPushButton:pressed {
    background-color: #d7e6f7;
}

QPushButton:disabled {
    background-color: #eceff2;
    border: 1px solid #d7dce3;
    color: #9aa4b2;
}

QPushButton#primaryButton {
    background-color: #1f4e79;
    border: 1px solid #1f4e79;
    color: #ffffff;
}

QPushButton#primaryButton:hover {
    background-color: #2c65a0;
}

QPushButton#primaryButton:pressed {
    background-color: #163a5c;
}

QPushButton#primaryButton:disabled {
    background-color: #eceff2;
    border: 1px solid #d7dce3;
    color: #9aa4b2;
}

QPushButton#dangerButton {
    background-color: #ffffff;
    border: 1px solid #c94b4b;
    color: #c94b4b;
}

QPushButton#dangerButton:hover {
    background-color: #fbeaea;
}

QPushButton#tableActionButton {
    background-color: #ffffff;
    border: 1px solid #1f4e79;
    color: #1f4e79;
    border-radius: 5px;
    padding: 5px 14px;
    font-weight: 600;
    font-size: 9pt;
}

QPushButton#tableActionButton:hover {
    background-color: #1f4e79;
    color: #ffffff;
}

QPushButton#tableActionButton:pressed {
    background-color: #163a5c;
    border: 1px solid #163a5c;
    color: #ffffff;
}

QPushButton#zoomButton {
    padding: 0px;
    font-size: 14pt;
    font-weight: 700;
}

QLabel {
    color: #1c2530;
    background-color: transparent;
}

QScrollBar:vertical {
    background: #f3f5f8;
    width: 12px;
    margin: 0;
}

QScrollBar::handle:vertical {
    background: #c3ccd8;
    border-radius: 6px;
    min-height: 24px;
}

QScrollBar::handle:vertical:hover {
    background: #9aa7b8;
}

QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {
    height: 0;
}
"""



class LoginDialog(QDialog):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("CPP - Login")
        self.setFixedWidth(340)
        self.user = None

        self.username_input = QLineEdit("")
        self.password_input = QLineEdit()
        self.password_input.setEchoMode(QLineEdit.Password)
        self.password_input.setText("")

        login_button = QPushButton("Entrar")
        login_button.clicked.connect(self.try_login)

        layout = QFormLayout(self)
        layout.addRow("Usuario", self.username_input)
        layout.addRow("Senha", self.password_input)
        layout.addRow(login_button)

    def try_login(self):
        user = authenticate(self.username_input.text(), self.password_input.text())
        if not user:
            QMessageBox.warning(self, "Acesso negado", "Usuario ou senha invalidos.")
            return
        self.user = user
        self.accept()


class PrintPreviewDialog(QDialog):
    """Pre-visualizacao e impressao 100% dentro do app, sem abrir telas do Windows."""

    def __init__(self, pdf_document, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Pre-visualizacao de impressao")
        self.resize(900, 700)

        self.pdf_document = pdf_document
        self.printer = QPrinter(QPrinter.PrinterMode.HighResolution)

        default_printer = QPrinterInfo.defaultPrinter()
        if not default_printer.isNull():
            self.printer.setPrinterName(default_printer.printerName())

        self.printer_combo = QComboBox()
        for printer_info in QPrinterInfo.availablePrinters():
            self.printer_combo.addItem(printer_info.printerName())

        current_name = self.printer.printerName()
        index = self.printer_combo.findText(current_name)
        if index >= 0:
            self.printer_combo.setCurrentIndex(index)

        self.printer_combo.currentTextChanged.connect(self._change_printer)

        self.preview_widget = QPrintPreviewWidget(self.printer, self)
        self.preview_widget.paintRequested.connect(self._render_pages)
        self.preview_widget.setLandscapeOrientation()

        self.orientation_combo = QComboBox()
        self.orientation_combo.addItem("Paisagem (horizontal)", "landscape")
        self.orientation_combo.addItem("Retrato (vertical)", "portrait")
        self.orientation_combo.currentIndexChanged.connect(self._change_orientation)

        zoom_out_button = QPushButton("-")
        zoom_out_button.setObjectName("zoomButton")
        zoom_out_button.setFixedSize(30, 30)
        zoom_out_button.clicked.connect(self._zoom_out)

        zoom_in_button = QPushButton("+")
        zoom_in_button.setObjectName("zoomButton")
        zoom_in_button.setFixedSize(30, 30)
        zoom_in_button.clicked.connect(self._zoom_in)

        self.zoom_label = QLabel("100%")
        self.zoom_label.setFixedWidth(45)
        self.zoom_label.setAlignment(Qt.AlignCenter)

        fit_width_button = QPushButton("Ajustar largura")
        fit_width_button.clicked.connect(self._fit_width)

        fit_page_button = QPushButton("Ajustar pagina")
        fit_page_button.clicked.connect(self._fit_page)

        print_button = QPushButton("Imprimir")
        print_button.setObjectName("primaryButton")
        print_button.clicked.connect(self._print_now)

        close_button = QPushButton("Fechar")
        close_button.clicked.connect(self.reject)

        top_layout = QHBoxLayout()
        top_layout.addWidget(QLabel("Impressora"))
        top_layout.addWidget(self.printer_combo)
        top_layout.addWidget(QLabel("Orientacao"))
        top_layout.addWidget(self.orientation_combo)
        top_layout.addStretch()
        top_layout.addWidget(QLabel("Zoom"))
        top_layout.addWidget(zoom_out_button)
        top_layout.addWidget(self.zoom_label)
        top_layout.addWidget(zoom_in_button)
        top_layout.addWidget(fit_width_button)
        top_layout.addWidget(fit_page_button)

        bottom_layout = QHBoxLayout()
        bottom_layout.addStretch()
        bottom_layout.addWidget(print_button)
        bottom_layout.addWidget(close_button)

        layout = QVBoxLayout(self)
        layout.addLayout(top_layout)
        layout.addWidget(self.preview_widget)
        layout.addLayout(bottom_layout)

        self._update_zoom_label()

    def _zoom_in(self):
        self.preview_widget.zoomIn()
        self._update_zoom_label()

    def _zoom_out(self):
        self.preview_widget.zoomOut()
        self._update_zoom_label()

    def _fit_width(self):
        self.preview_widget.fitToWidth()
        self._update_zoom_label()

    def _fit_page(self):
        self.preview_widget.fitInView()
        self._update_zoom_label()

    def _change_orientation(self):
        if self.orientation_combo.currentData() == "landscape":
            self.preview_widget.setLandscapeOrientation()
        else:
            self.preview_widget.setPortraitOrientation()
        self._update_zoom_label()

    def _update_zoom_label(self):
        percent = round(self.preview_widget.zoomFactor() * 100)
        self.zoom_label.setText(f"{percent}%")

    def _change_printer(self, printer_name):
        if printer_name:
            self.printer.setPrinterName(printer_name)
        self.preview_widget.updatePreview()

    def _render_pages(self, printer):
        painter = QPainter(printer)
        try:
            for page in range(self.pdf_document.pageCount()):
                if page > 0:
                    printer.newPage()

                page_size = self.pdf_document.pagePointSize(page).toSize() * 4
                image = self.pdf_document.render(page, page_size)

                target = painter.viewport()
                scaled = image.scaled(
                    target.size(),
                    Qt.AspectRatioMode.KeepAspectRatio,
                    Qt.TransformationMode.SmoothTransformation,
                )
                x = target.x() + (target.width() - scaled.width()) // 2
                y = target.y() + (target.height() - scaled.height()) // 2
                painter.drawImage(x, y, scaled)
        finally:
            painter.end()

    def _print_now(self):
        self._render_pages(self.printer)
        self.accept()


def _normalize_lookup_text(value) -> str:
    import re
    import unicodedata
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).upper()


def _looks_like_price_header_row(row) -> bool:
    values = [_normalize_lookup_text(value) for value in row]
    return "PROCEDIMENTO" in values[:6] and "DESCRICAO" in values[:6]


def _extract_doctors_and_organizations_from_price_table(path):
    """Lê os nomes de médicos/organizações dos blocos da planilha.

    A planilha organiza os procedimentos em blocos: o nome do prestador
    aparece antes da linha de cabeçalho "Procedimento". A chave também leva
    em conta os três valores da tabela para não misturar dois prestadores que
    tenham o mesmo código/descrição, mas preços diferentes.

    Retorna:
      - providers_by_key: (aba, código, descrição, valores) -> conjunto de prestadores
      - providers_by_source_code: (aba, código) -> conjunto de prestadores
      - providers_by_code: código -> conjunto de prestadores
    """
    providers_by_key = {}
    providers_by_source_code = {}
    providers_by_code = {}

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return providers_by_key, providers_by_source_code, providers_by_code

    import re

    def money_key(value):
        if value in (None, ""):
            return None
        try:
            return round(float(value), 6)
        except (TypeError, ValueError):
            return None

    for sheet in workbook.worksheets:
        pending_label = ""
        current_provider = ""

        for row in sheet.iter_rows(values_only=True):
            values = list(row)
            non_empty = [value for value in values if value not in (None, "")]

            if _looks_like_price_header_row(values):
                if pending_label:
                    current_provider = pending_label
                pending_label = ""
                continue

            # Linha com apenas um texto: nome do médico/clínica/organização
            # que inicia o próximo bloco.
            if len(non_empty) == 1 and isinstance(non_empty[0], str):
                label = non_empty[0].strip()
                if label and not re.fullmatch(r"[-=]+", label):
                    pending_label = label

            if not values:
                continue

            raw_code = values[0]
            if raw_code is None:
                continue

            code_text = str(raw_code).strip()
            code_digits = re.sub(r"\D", "", code_text)
            if len(code_digits) < 6:
                continue

            description = values[1] if len(values) > 1 else None
            if not isinstance(description, str) or not description.strip():
                continue

            # Na planilha da imagem, os preços ficam nas colunas D/E/F
            # (PARTICULAR/PARCEIROS/PRÓ-SAÚDE PREMIUM). Usamos os valores
            # reais da linha para diferenciar blocos com a mesma descrição.
            price_signature = (
                money_key(values[3] if len(values) > 3 else None),
                money_key(values[4] if len(values) > 4 else None),
                money_key(values[5] if len(values) > 5 else None),
            )

            provider = (current_provider or sheet.title).strip()

            key = (
                sheet.title,
                code_digits,
                _normalize_lookup_text(description),
                price_signature,
            )

            providers_by_key.setdefault(key, set()).add(provider)
            providers_by_source_code.setdefault(
                (sheet.title, code_digits), set()
            ).add(provider)
            providers_by_code.setdefault(code_digits, set()).add(provider)

    workbook.close()
    return providers_by_key, providers_by_source_code, providers_by_code


class MainWindow(QMainWindow):
    def __init__(self, user):
        super().__init__()
        self.user = user
        self.logged_out = False
        self.current_rows = []
        self.current_preview_kind = "report"
        self.current_procedure = None
        self.pending_exams = []
        self.editing_attendance_id = None
        self.editing_mode = False
        self.editing_patient_name = None
        self.original_exam_ids = set()
        self.editing_pending_exam_id = None
        self.current_budget_procedure = None
        self.budget_items = []
        self.price_table_provider_by_key, self.price_table_provider_by_source_code, self.price_table_providers_by_code = _extract_doctors_and_organizations_from_price_table(PRICE_TABLE_PATH)
        self.setWindowTitle(f"CPP - Controle de Passagem de Paciente ({user.username})")
        self.resize(1080, 720)

        self.report_preview_table = QTableWidget()

        root = QWidget()
        main_layout = QVBoxLayout(root)

        top_bar = QHBoxLayout()
        top_bar.addStretch()

        logout_button = QPushButton("Sair")
        logout_button.setObjectName("dangerButton")
        logout_button.clicked.connect(self.logout)
        top_bar.addWidget(logout_button)

        main_layout.addLayout(top_bar)

        self.tabs = QTabWidget()
        tabs = self.tabs

        # Aba 1
        attendance_tab = QWidget()
        self.attendance_tab = attendance_tab
        attendance_layout = QVBoxLayout(attendance_tab)
        attendance_layout.addWidget(self.build_attendance_group())

        # Aba 2
        search_tab = QWidget()
        search_layout = QVBoxLayout(search_tab)
        search_layout.addWidget(self.build_search_group())
        search_layout.addWidget(self.build_table())

        # Aba 3
        reports_tab = QWidget()
        reports_layout = QVBoxLayout(reports_tab)
        reports_layout.addWidget(self.build_reports_group())

        # Aba 4
        budget_tab = QWidget()
        budget_layout = QVBoxLayout(budget_tab)
        budget_layout.addWidget(self.build_budget_group())

        tabs.addTab(attendance_tab, "Atendimentos")
        tabs.addTab(search_tab, "Pesquisa")
        tabs.addTab(reports_tab, "Relatórios")
        tabs.addTab(budget_tab, "Orçamento")

        # Aba 5 - protegida por senha, so libera uso se for admin
        tabs.addTab(self.build_users_tab(), "Usuarios")

        main_layout.addWidget(tabs)


        self.setCentralWidget(root)

        self.refresh_table()

        self.load_report_preview()

    def logout(self):
        confirm = QMessageBox.question(
            self,
            "Sair",
            "Tem certeza que deseja sair do sistema?",
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        self.logged_out = True
        self.close()

    def build_attendance_group(self) -> QGroupBox:
        group = QGroupBox("Cadastro de atendimento")
        main_layout = QHBoxLayout(group)

        left_column = QWidget()
        layout = QFormLayout(left_column)

        self.patient_input = QLineEdit()
        self.patient_input.setAlignment(Qt.AlignCenter)
        patient_font = self.patient_input.font()
        patient_font.setPointSize(patient_font.pointSize() + 2)
        patient_font.setBold(True)
        self.patient_input.setFont(patient_font)
        self.patient_input.setPlaceholderText("Nome do paciente")

        self.procedure_search_input = QLineEdit()
        self.procedure_search_input.setPlaceholderText(
            "Digite código ou nome..."
        )

        self.procedure_search_input.textChanged.connect(
            self.search_procedures
        )

        self.procedure_list = QListWidget()
        self.procedure_list.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        self.procedure_list.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.procedure_list.setTextElideMode(Qt.TextElideMode.ElideNone)

        self.procedure_list.itemDoubleClicked.connect(
            self.select_procedure
)
        self.procedure_name_label = QLabel("-")
        self.particular_value_label = QLabel("-")
        self.partner_value_label = QLabel("-")
        self.premium_value_label = QLabel("-")
        self.selected_price_label = QLabel("-")
        self.procedure_source_label = QLabel("-")

        self.attendance_date_input = QLineEdit()
        self.attendance_date_input.setPlaceholderText("dd/mm/aaaa")
        self.attendance_date_input.setText(QDate.currentDate().toString("dd/MM/yyyy"))
        self.attendance_date_input.setMaxLength(10)
        self.attendance_date_input.textEdited.connect(
            lambda text: self._auto_format_date_text(self.attendance_date_input, text)
        )
        self.attendance_date_input.editingFinished.connect(
            lambda: self._finish_date_input(
                self.attendance_date_input,
                on_valid=lambda parsed_date: self.sync_report_month_year(
                    QDate(parsed_date.year, parsed_date.month, parsed_date.day)
                ),
            )
        )

        self.paid_value_input = QLineEdit()
        self.sps_code_input = QLineEdit()

        self.price_type_combo = QComboBox()
        self.price_type_combo.addItems(
            ["PARTICULAR", "PARCEIROS", "PRO PREMIUM"]
        )
        self.price_type_combo.currentTextChanged.connect(
            self.update_selected_price
        )

        self.report_month_combo = QComboBox()

        for number, name in MONTH_NAMES.items():
            self.report_month_combo.addItem(name, number)

        self.report_month_combo.setCurrentIndex(
            QDate.currentDate().month() - 1
        )

        self.report_year_combo = QComboBox()

        current_year = QDate.currentDate().year()

        for year in range(current_year - 1, current_year + 3):
            self.report_year_combo.addItem(str(year), year)

        self.report_year_combo.setCurrentText(str(current_year))

        self.add_exam_button = QPushButton("Adicionar exame")
        self.add_exam_button.setObjectName("primaryButton")
        self.add_exam_button.clicked.connect(self.add_exam_to_patient)

        self.exams_table = QTableWidget()
        self.exams_table.setColumnCount(5)
        self.exams_table.setHorizontalHeaderLabels(
            ["Procedimento", "Tipo", "Valor pago", "SAVE", "Data"]
        )
        self.exams_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.exams_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.exams_table.setAlternatingRowColors(True)
        self.exams_table.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.exams_table.setMinimumHeight(260)
        self.exams_table.horizontalHeader().setStretchLastSection(True)
        self.exams_table.cellDoubleClicked.connect(self.edit_pending_exam)

        remove_exam_button = QPushButton("Remover exame selecionado")
        remove_exam_button.setObjectName("dangerButton")
        remove_exam_button.clicked.connect(self.remove_selected_exam)

        self.save_button = QPushButton("Salvar atendimentos do paciente")
        self.save_button.setObjectName("primaryButton")
        self.save_button.clicked.connect(self.save_patient_attendances)

        self.cancel_edit_button = QPushButton("Cancelar edicao")
        self.cancel_edit_button.setObjectName("dangerButton")
        self.cancel_edit_button.clicked.connect(self.cancel_edit_attendance)
        self.cancel_edit_button.setVisible(False)

        self.edit_mode_label = QLabel("")
        self.edit_mode_label.setWordWrap(True)
        edit_mode_font = self.edit_mode_label.font()
        edit_mode_font.setBold(True)
        self.edit_mode_label.setFont(edit_mode_font)
        self.edit_mode_label.setStyleSheet("color: #b30000;")
        self.edit_mode_label.setVisible(False)

        self.update_price_table_button = QPushButton("Atualizar tabela de preços")
        self.update_price_table_button.setObjectName("primaryButton")
        self.update_price_table_button.clicked.connect(self.update_price_table)

        layout.addRow(self.update_price_table_button)

        layout.addRow(self.edit_mode_label)

        layout.addRow("Paciente", self.patient_input)

        layout.addRow(
            "Pesquisar",
            self.procedure_search_input
)

        layout.addRow(
            "",
            self.procedure_list
)

        layout.addRow("Procedimento", self.procedure_name_label)
        layout.addRow("Origem tabela", self.procedure_source_label)

        layout.addRow("Particular", self.particular_value_label)
        layout.addRow("Parceiros", self.partner_value_label)
        layout.addRow("Pro Premium", self.premium_value_label)

        layout.addRow("Tipo de preco", self.price_type_combo)
        layout.addRow("Valor cobrado", self.selected_price_label)

        layout.addRow("Data atendimento", self.attendance_date_input)

        month_layout = QHBoxLayout()
        month_layout.addWidget(self.report_month_combo)
        month_layout.addWidget(self.report_year_combo)

        layout.addRow("Mes relatorio", month_layout)

        layout.addRow("Valor pago", self.paid_value_input)
        layout.addRow("SAVE", self.sps_code_input)

        layout.addRow(self.add_exam_button)

        layout.setVerticalSpacing(6)
        layout.setHorizontalSpacing(15)

        right_column = QWidget()
        right_layout = QVBoxLayout(right_column)

        exams_label = QLabel("Exames deste paciente")
        exams_label_font = exams_label.font()
        exams_label_font.setBold(True)
        exams_label_font.setPointSize(exams_label_font.pointSize() + 1)
        exams_label.setFont(exams_label_font)

        right_layout.addWidget(exams_label)
        right_layout.addWidget(self.exams_table, 1)
        right_layout.addWidget(remove_exam_button)
        right_layout.addWidget(self.save_button)
        right_layout.addWidget(self.cancel_edit_button)

        main_layout.addWidget(left_column, 1)
        main_layout.addWidget(right_column, 1)

        return group

    def update_price_table(self):
        if not Path(PRICE_TABLE_PATH).exists():
            QMessageBox.warning(
                self,
                "Tabela não encontrada",
                "Não foi possível localizar a tabela de preços no caminho configurado:\n"
                f"{PRICE_TABLE_PATH}",
            )
            return

        confirm = QMessageBox.question(
            self,
            "Atualizar tabela de preços",
            "Atualizar a tabela de preços a partir do arquivo configurado?\n\n"
            f"{PRICE_TABLE_PATH}",
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        try:
            refresh_procedure_prices()
        except Exception as exc:
            QMessageBox.critical(
                self,
                "Erro ao atualizar",
                f"Ocorreu um erro ao atualizar a tabela de preços:\n{exc}",
            )
            return

        # Recarrega os prestadores/organizacoes extraidos da planilha e
        # atualiza a lista de procedimentos exibida, se houver pesquisa ativa.
        (
            self.price_table_provider_by_key,
            self.price_table_provider_by_source_code,
            self.price_table_providers_by_code,
        ) = _extract_doctors_and_organizations_from_price_table(PRICE_TABLE_PATH)

        self.search_procedures()

        QMessageBox.information(
            self,
            "Tabela atualizada",
            "Tabela de preços atualizada com sucesso.",
        )

    def build_search_group(self) -> QGroupBox:
        group = QGroupBox("Consulta de registros")
        layout = QHBoxLayout(group)

        self.search_patient_input = QLineEdit()
        self.search_patient_input.setPlaceholderText("Paciente")

        self.search_procedure_input = QLineEdit()
        self.search_procedure_input.setPlaceholderText("Procedimento")

        self.search_sps_code_input = QLineEdit()
        self.search_sps_code_input.setPlaceholderText("SAVE")

        self.search_date_mode = QComboBox()
        self.search_date_mode.addItems(
            ["Todas as datas", "Somente data informada"]
        )

        self.search_date_input = QLineEdit()
        self.search_date_input.setPlaceholderText("dd/mm/aaaa")
        self.search_date_input.setText(QDate.currentDate().toString("dd/MM/yyyy"))
        self.search_date_input.setMaxLength(10)
        self.search_date_input.textEdited.connect(
            lambda text: self._auto_format_date_text(self.search_date_input, text)
        )
        self.search_date_input.editingFinished.connect(
            lambda: self._finish_date_input(self.search_date_input)
        )

        search_button = QPushButton("Pesquisar")
        search_button.setObjectName("primaryButton")
        search_button.clicked.connect(self.refresh_table)

        layout.addWidget(self.search_patient_input)
        layout.addWidget(self.search_procedure_input)
        layout.addWidget(self.search_sps_code_input)
        layout.addWidget(self.search_date_mode)
        layout.addWidget(self.search_date_input)
        layout.addWidget(search_button)

        return group
    
    def build_reports_group(self) -> QGroupBox:
        group = QGroupBox("Geração de relatórios")

        main_layout = QVBoxLayout(group)

        # ===== FILTROS =====

        top_layout = QHBoxLayout()

        self.export_month_start_combo = QComboBox()
        self.export_month_end_combo = QComboBox()

        for number, name in MONTH_NAMES.items():
            self.export_month_start_combo.addItem(name, number)
            self.export_month_end_combo.addItem(name, number)

        self.export_month_start_combo.setCurrentIndex(QDate.currentDate().month() - 1)
        self.export_month_end_combo.setCurrentIndex(QDate.currentDate().month() - 1)

        self.export_year_start_combo = QComboBox()
        self.export_year_end_combo = QComboBox()

        current_year = QDate.currentDate().year()

        for year in range(current_year - 2, current_year + 5):
            self.export_year_start_combo.addItem(str(year), year)
            self.export_year_end_combo.addItem(str(year), year)

        self.export_year_start_combo.setCurrentText(str(current_year))
        self.export_year_end_combo.setCurrentText(str(current_year))

        excel_button = QPushButton("Gerar Excel")
        excel_button.clicked.connect(self.export_current_excel)

        pdf_button = QPushButton("Gerar PDF")
        pdf_button.clicked.connect(self.export_current_pdf)

        print_button = QPushButton("Imprimir")
        print_button.clicked.connect(self.print_current_report)

        refresh_button = QPushButton("Atualizar")
        refresh_button.clicked.connect(self.load_report_preview)

        view_budget_button = QPushButton("Ver Orçamento")
        view_budget_button.clicked.connect(self.show_budget_preview)

        top_layout.addWidget(QLabel("Do mes"))
        top_layout.addWidget(self.export_month_start_combo)
        top_layout.addWidget(self.export_year_start_combo)

        top_layout.addWidget(QLabel("ate o mes"))
        top_layout.addWidget(self.export_month_end_combo)
        top_layout.addWidget(self.export_year_end_combo)

        top_layout.addWidget(excel_button)
        top_layout.addWidget(pdf_button)
        top_layout.addWidget(print_button)
        top_layout.addWidget(refresh_button)
        top_layout.addWidget(view_budget_button)

        top_layout.addStretch()

        main_layout.addLayout(top_layout)

        # ===== PREVIEW =====

        self.pdf_document = QPdfDocument(self)

        self.pdf_view = QPdfView()
        self.pdf_view.setDocument(self.pdf_document)

        main_layout.addWidget(self.pdf_view)

        self.report_preview_table.horizontalHeader().setStretchLastSection(True)

        

        # ===== EVENTOS =====

        self.export_month_start_combo.currentIndexChanged.connect(
            self.load_report_preview
        )

        self.export_month_end_combo.currentIndexChanged.connect(
            self.load_report_preview
        )

        self.export_year_start_combo.currentIndexChanged.connect(
            self.load_report_preview
        )

        self.export_year_end_combo.currentIndexChanged.connect(
            self.load_report_preview
        )

        self.load_report_preview()

        return group

    def _rows_for_range(self, mes_inicial: int, ano_inicial: int, mes_final: int, ano_final: int) -> list:
        inicio = ano_inicial * 100 + mes_inicial
        fim = ano_final * 100 + mes_final

        filtradas = [
            item
            for item in self.current_rows
            if inicio
            <= (
                (item.report_year or item.attendance_date.year) * 100
                + (item.report_month or item.attendance_date.month)
            )
            <= fim
        ]

        filtradas.sort(
            key=lambda item: (
                item.report_year or item.attendance_date.year,
                item.report_month or item.attendance_date.month,
                item.attendance_date,
            )
        )

        return filtradas

    def _resolve_report_range(self):
        """Le os combos de intervalo, valida e devolve (mes_inicial, ano_inicial, mes_final, ano_final, rows).
        Mostra os avisos necessarios e devolve None se nao der pra prosseguir."""
        mes_inicial = self.export_month_start_combo.currentData()
        ano_inicial = self.export_year_start_combo.currentData()
        mes_final = self.export_month_end_combo.currentData()
        ano_final = self.export_year_end_combo.currentData()

        if (ano_inicial, mes_inicial) > (ano_final, mes_final):
            QMessageBox.warning(
                self,
                "Intervalo invalido",
                "A data inicial nao pode ser depois da data final."
            )
            return None

        rows = self._rows_for_range(mes_inicial, ano_inicial, mes_final, ano_final)

        if not rows:
            QMessageBox.information(
                self,
                "Relatorio",
                f"Nao existem registros entre {MONTH_NAMES[mes_inicial]} {ano_inicial} "
                f"e {MONTH_NAMES[mes_final]} {ano_final}."
            )
            return None

        return mes_inicial, ano_inicial, mes_final, ano_final, rows

    def _report_range_filename(
        self, mes_inicial: int, ano_inicial: int, mes_final: int, ano_final: int, extension: str
    ) -> str:
        if mes_inicial == mes_final and ano_inicial == ano_final:
            return f"Relatorio_{MONTH_NAMES[mes_inicial]}_{ano_inicial}.{extension}"
        return (
            f"Relatorio_{MONTH_NAMES[mes_inicial]}_{ano_inicial}"
            f"_a_{MONTH_NAMES[mes_final]}_{ano_final}.{extension}"
        )

    def load_report_preview(self):
        mes_inicial = self.export_month_start_combo.currentData()
        ano_inicial = self.export_year_start_combo.currentData()
        mes_final = self.export_month_end_combo.currentData()
        ano_final = self.export_year_end_combo.currentData()

        if (ano_inicial, mes_inicial) > (ano_final, mes_final):
            return

        rows = self._rows_for_range(mes_inicial, ano_inicial, mes_final, ano_final)

        if not rows:
            return

        temp_pdf = Path(tempfile.gettempdir()) / "cpp_preview.pdf"

        export_pdf(
            rows,
            filename=str(temp_pdf)
        )

        self.current_preview_kind = "report"
        self.pdf_document.load(str(temp_pdf))

    def build_budget_group(self) -> QGroupBox:
        group = QGroupBox("Orçamento")
        main_layout = QHBoxLayout(group)

        left_column = QWidget()
        form_layout = QFormLayout(left_column)

        self.budget_search_input = QLineEdit()
        self.budget_search_input.setPlaceholderText("Digite código ou nome...")
        self.budget_search_input.textChanged.connect(self.search_budget_procedures)

        self.budget_procedure_list = QListWidget()
        self.budget_procedure_list.itemDoubleClicked.connect(self.select_budget_procedure)

        self.budget_code_label = QLabel("-")
        self.budget_procedure_name_label = QLabel("-")
        self.budget_particular_label = QLabel("-")
        self.budget_partner_label = QLabel("-")
        self.budget_premium_label = QLabel("-")

        add_budget_button = QPushButton("Adicionar procedimento")
        add_budget_button.setObjectName("primaryButton")
        add_budget_button.clicked.connect(self.add_procedure_to_budget)

        form_layout.addRow("Pesquisar", self.budget_search_input)
        form_layout.addRow("", self.budget_procedure_list)
        form_layout.addRow("Codigo", self.budget_code_label)
        form_layout.addRow("Procedimento", self.budget_procedure_name_label)
        form_layout.addRow("Particular", self.budget_particular_label)
        form_layout.addRow("Parceiros", self.budget_partner_label)
        form_layout.addRow("Pro Premium", self.budget_premium_label)
        form_layout.addRow(add_budget_button)

        form_layout.setVerticalSpacing(6)
        form_layout.setHorizontalSpacing(15)

        right_column = QWidget()
        right_layout = QVBoxLayout(right_column)

        budget_label = QLabel("Procedimentos do orçamento")
        budget_label_font = budget_label.font()
        budget_label_font.setBold(True)
        budget_label_font.setPointSize(budget_label_font.pointSize() + 1)
        budget_label.setFont(budget_label_font)

        self.budget_table = QTableWidget()
        self.budget_table.setColumnCount(5)
        self.budget_table.setHorizontalHeaderLabels(
            ["Codigo", "Procedimento", "Particular", "Parceiros", "Pro Premium"]
        )
        self.budget_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.budget_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.budget_table.setAlternatingRowColors(True)
        self.budget_table.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.budget_table.setMinimumHeight(260)
        self.budget_table.horizontalHeader().setStretchLastSection(True)

        remove_budget_button = QPushButton("Remover procedimento selecionado")
        remove_budget_button.setObjectName("dangerButton")
        remove_budget_button.clicked.connect(self.remove_selected_budget_item)

        totals_group = QGroupBox("Totais do orçamento")
        totals_layout = QHBoxLayout(totals_group)

        self.budget_total_particular_label = QLabel("R$ 0,00")
        self.budget_total_partner_label = QLabel("R$ 0,00")
        self.budget_total_premium_label = QLabel("R$ 0,00")

        for total_label in (
            self.budget_total_particular_label,
            self.budget_total_partner_label,
            self.budget_total_premium_label,
        ):
            total_font = total_label.font()
            total_font.setBold(True)
            total_label.setFont(total_font)

        particular_box = QVBoxLayout()
        particular_box.addWidget(QLabel("Particular"))
        particular_box.addWidget(self.budget_total_particular_label)

        partner_box = QVBoxLayout()
        partner_box.addWidget(QLabel("Parceiros"))
        partner_box.addWidget(self.budget_total_partner_label)

        premium_box = QVBoxLayout()
        premium_box.addWidget(QLabel("Pro Premium"))
        premium_box.addWidget(self.budget_total_premium_label)

        totals_layout.addLayout(particular_box)
        totals_layout.addLayout(partner_box)
        totals_layout.addLayout(premium_box)

        new_budget_button = QPushButton("Gerar novo orçamento")
        new_budget_button.setObjectName("primaryButton")
        new_budget_button.clicked.connect(self.start_new_budget)

        right_layout.addWidget(budget_label)
        right_layout.addWidget(self.budget_table, 1)
        right_layout.addWidget(totals_group)
        right_layout.addWidget(remove_budget_button)
        right_layout.addWidget(new_budget_button)

        main_layout.addWidget(left_column, 1)
        main_layout.addWidget(right_column, 1)

        return group

    def search_budget_procedures(self):
        texto = self.budget_search_input.text().strip()

        self.budget_procedure_list.clear()

        if len(texto) < 2:
            return

        with SessionLocal() as db:
            resultados = (
                db.query(Procedure)
                .filter(
                    or_(
                        Procedure.name.ilike(f"%{texto}%"),
                        Procedure.code.ilike(f"%{texto}%")
                    )
                )
                .order_by(Procedure.name)
                .all()
            )

        for proc in resultados:
            codigo = proc.code if proc.code else "SEM CÓDIGO"

            item = QListWidgetItem(f"{codigo} - {proc.name}")
            item.setData(Qt.UserRole, proc)

            self.budget_procedure_list.addItem(item)

    def select_budget_procedure(self, item):
        procedure = item.data(Qt.UserRole)
        self.current_budget_procedure = procedure
        self.budget_procedure_list.clear()

        codigo = procedure.code if procedure.code else "SEM CÓDIGO"
        self.budget_code_label.setText(codigo)
        self.budget_procedure_name_label.setText(procedure.name)
        self.budget_particular_label.setText(self.format_money(procedure.reference_value))
        self.budget_partner_label.setText(self.format_money(procedure.partner_value))
        self.budget_premium_label.setText(self.format_money(procedure.premium_value))

    def add_procedure_to_budget(self):
        if not self.current_budget_procedure:
            QMessageBox.warning(self, "Procedimento", "Selecione um procedimento na lista antes de adicionar.")
            return

        self.budget_items.append(self.current_budget_procedure)
        self._refresh_budget_table()

        self.current_budget_procedure = None
        self.budget_search_input.clear()
        self.budget_code_label.setText("-")
        self.budget_procedure_name_label.setText("-")
        self.budget_particular_label.setText("-")
        self.budget_partner_label.setText("-")
        self.budget_premium_label.setText("-")

    def _refresh_budget_table(self):
        self.budget_table.setRowCount(len(self.budget_items))

        total_particular = 0.0
        total_partner = 0.0
        total_premium = 0.0

        for row, procedure in enumerate(self.budget_items):
            codigo = procedure.code if procedure.code else "SEM CÓDIGO"
            values = [
                codigo,
                procedure.name,
                self.format_money(procedure.reference_value),
                self.format_money(procedure.partner_value),
                self.format_money(procedure.premium_value),
            ]
            for column, value in enumerate(values):
                self.budget_table.setItem(row, column, QTableWidgetItem(value))

            total_particular += float(procedure.reference_value or 0)
            total_partner += float(procedure.partner_value or 0)
            total_premium += float(procedure.premium_value or 0)

        self.budget_table.resizeColumnsToContents()

        self.budget_total_particular_label.setText(self.format_money(total_particular))
        self.budget_total_partner_label.setText(self.format_money(total_partner))
        self.budget_total_premium_label.setText(self.format_money(total_premium))

    def remove_selected_budget_item(self):
        row = self.budget_table.currentRow()
        if row < 0:
            return
        del self.budget_items[row]
        self._refresh_budget_table()

    def start_new_budget(self):
        if not self.budget_items:
            return

        confirm = QMessageBox.question(
            self,
            "Gerar novo orçamento",
            "Isso vai remover todos os procedimentos adicionados agora, pra você comecar um orcamento novo. Deseja continuar?",
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        self.budget_items = []
        self._refresh_budget_table()

    def show_budget_preview(self):
        if not self.budget_items:
            QMessageBox.information(
                self,
                "Orçamento",
                "Adicione ao menos um procedimento na aba Orçamento antes de visualizar."
            )
            return

        temp_pdf = Path(tempfile.gettempdir()) / "cpp_budget_preview.pdf"
        export_budget_pdf(self.budget_items, filename=str(temp_pdf))

        self.pdf_document.load(str(temp_pdf))
        self.current_preview_kind = "budget"

    def build_users_tab(self) -> QWidget:
        tab = QWidget()
        tab_layout = QVBoxLayout(tab)

        self.users_stack = QStackedWidget()

        self.users_stack.addWidget(self._build_users_lock_screen())
        self.users_stack.addWidget(self.build_users_group())
        self.users_stack.setCurrentIndex(0)

        tab_layout.addWidget(self.users_stack)

        return tab

    def _build_users_lock_screen(self) -> QGroupBox:
        group = QGroupBox("Area restrita")
        layout = QFormLayout(group)

        info_label = QLabel("Digite usuario e senha de administrador para acessar o cadastro de usuarios.")
        info_label.setWordWrap(True)

        self.users_unlock_username_input = QLineEdit()
        self.users_unlock_username_input.setPlaceholderText("Usuario")

        self.users_unlock_password_input = QLineEdit()
        self.users_unlock_password_input.setPlaceholderText("Senha")
        self.users_unlock_password_input.setEchoMode(QLineEdit.EchoMode.Password)
        self.users_unlock_password_input.returnPressed.connect(self.try_unlock_users_tab)

        self.users_unlock_username_input.returnPressed.connect(self.users_unlock_password_input.setFocus)

        unlock_button = QPushButton("Entrar")
        unlock_button.setObjectName("primaryButton")
        unlock_button.clicked.connect(self.try_unlock_users_tab)

        layout.addRow(info_label)
        layout.addRow("Usuario", self.users_unlock_username_input)
        layout.addRow("Senha", self.users_unlock_password_input)
        layout.addRow(unlock_button)

        return group

    def try_unlock_users_tab(self):
        username = self.users_unlock_username_input.text()
        password = self.users_unlock_password_input.text()

        checked_user = authenticate(username, password)

        if not checked_user:
            QMessageBox.warning(self, "Acesso negado", "Usuario ou senha invalidos.")
            self.users_unlock_password_input.clear()
            return

        if checked_user.role != "admin":
            QMessageBox.warning(self, "Acesso negado", "Apenas administradores podem acessar esta area.")
            self.users_unlock_password_input.clear()
            return

        self.users_unlock_username_input.clear()
        self.users_unlock_password_input.clear()
        self.refresh_users_table()
        self.users_stack.setCurrentIndex(1)

    def build_users_group(self) -> QGroupBox:
        group = QGroupBox("Cadastro de usuarios")
        layout = QFormLayout(group)

        self.new_username_input = QLineEdit()
        self.new_username_input.setPlaceholderText("Nome de usuario")

        self.new_password_input = QLineEdit()
        self.new_password_input.setPlaceholderText("Senha")
        self.new_password_input.setEchoMode(QLineEdit.EchoMode.Password)

        self.new_role_combo = QComboBox()
        self.new_role_combo.addItem("Usuario", "user")
        self.new_role_combo.addItem("Administrador", "admin")

        save_user_button = QPushButton("Salvar usuario")
        save_user_button.setObjectName("primaryButton")
        save_user_button.clicked.connect(self.save_new_user)

        self.users_table = QTableWidget()
        self.users_table.setColumnCount(3)
        self.users_table.setHorizontalHeaderLabels(["Nome", "Classe", "Criado em"])
        self.users_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.users_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.users_table.setAlternatingRowColors(True)
        self.users_table.horizontalHeader().setStretchLastSection(True)

        delete_user_button = QPushButton("Excluir usuario selecionado")
        delete_user_button.setObjectName("dangerButton")
        delete_user_button.clicked.connect(self.delete_selected_user)

        layout.addRow("Nome", self.new_username_input)
        layout.addRow("Senha", self.new_password_input)
        layout.addRow("Classe", self.new_role_combo)
        layout.addRow(save_user_button)
        layout.addRow(QLabel("Usuarios cadastrados"))
        layout.addRow(self.users_table)
        layout.addRow(delete_user_button)

        self.refresh_users_table()

        return group

    def save_new_user(self):
        username = self.new_username_input.text()
        password = self.new_password_input.text()
        role = self.new_role_combo.currentData()

        try:
            create_user(username, password, role)
        except ValueError as exc:
            QMessageBox.warning(self, "Usuario", str(exc))
            return

        self.new_username_input.clear()
        self.new_password_input.clear()
        self.new_role_combo.setCurrentIndex(0)
        self.refresh_users_table()

        QMessageBox.information(self, "Usuario", f"Usuario '{username.strip()}' criado com sucesso.")

    def delete_selected_user(self):
        row = self.users_table.currentRow()
        if row < 0:
            QMessageBox.information(self, "Excluir usuario", "Selecione um usuario na tabela primeiro.")
            return

        name_item = self.users_table.item(row, 0)
        user_id = name_item.data(Qt.ItemDataRole.UserRole)
        username = name_item.text()

        if username == self.user.username:
            QMessageBox.warning(self, "Excluir usuario", "Voce nao pode excluir o proprio usuario enquanto estiver logado.")
            return

        confirm = QMessageBox.question(
            self,
            "Excluir usuario",
            f"Tem certeza que deseja excluir o usuario '{username}'? Essa acao nao pode ser desfeita.",
        )
        if confirm != QMessageBox.StandardButton.Yes:
            return

        try:
            delete_user(user_id)
        except ValueError as exc:
            QMessageBox.warning(self, "Excluir usuario", str(exc))
            return

        self.refresh_users_table()
        QMessageBox.information(self, "Excluir usuario", f"Usuario '{username}' excluido com sucesso.")

    def refresh_users_table(self):
        users = list_users()
        self.users_table.setRowCount(len(users))
        for row, user in enumerate(users):
            values = [
                user.username,
                "Administrador" if user.role == "admin" else "Usuario",
                user.created_at.strftime("%d/%m/%Y %H:%M") if user.created_at else "-",
            ]
            for column, value in enumerate(values):
                cell = QTableWidgetItem(value)
                if column == 0:
                    cell.setData(Qt.ItemDataRole.UserRole, user.id)
                self.users_table.setItem(row, column, cell)
        self.users_table.resizeColumnsToContents()

    def build_table(self) -> QTableWidget:
        self.table = QTableWidget(0, 12)
        self.table.setHorizontalHeaderLabels(
    [
        "Paciente",
        "Data",
        "Mes",
        "Codigo",
        "Procedimento",
        "Tipo",
        "Tabela Origem",
        "Valor cobrado",
        "Valor pago",
        "Diferenca",
        "SAVE",
        "Acoes",
    ]

    

    
)

        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setAlternatingRowColors(True)
        return self.table

    def _display_procedure(self, procedure):
        self.current_procedure = procedure
        self.procedure_list.clear()
        self.procedure_name_label.setText(procedure.name)
        self.procedure_source_label.setText(procedure.source_sheet)
        self.particular_value_label.setText(self.format_money(procedure.reference_value))
        self.partner_value_label.setText(self.format_money(procedure.partner_value))
        self.premium_value_label.setText(self.format_money(procedure.premium_value))
        self.update_selected_price()

    @staticmethod
    def format_money(value):
        if value is None:
            return "-"
        return f"R$ {float(value):.2f}"

    def selected_price_value(self):
        if not self.current_procedure:
            return None
        price_type = self.price_type_combo.currentText()
        if price_type == "PARCEIROS":
            return self.current_procedure.partner_value or self.current_procedure.reference_value
        if price_type == "PRO PREMIUM":
            return self.current_procedure.premium_value or self.current_procedure.reference_value
        return self.current_procedure.reference_value

    def update_selected_price(self):
        self.selected_price_label.setText(self.format_money(self.selected_price_value()))

    def sync_report_month_year(self, qdate):
        self.report_month_combo.setCurrentIndex(qdate.month() - 1)

        year = qdate.year()
        index = self.report_year_combo.findData(year)
        if index == -1:
            self.report_year_combo.addItem(str(year), year)
            index = self.report_year_combo.findData(year)
        self.report_year_combo.setCurrentIndex(index)

    @staticmethod
    def _parse_date_digits(digits: str) -> date:
        """Converte 'ddMMyy' (6 digitos) ou 'ddMMyyyy' (8 digitos) em date."""
        if len(digits) == 6:
            day, month, year = int(digits[0:2]), int(digits[2:4]), 2000 + int(digits[4:6])
        elif len(digits) == 8:
            day, month, year = int(digits[0:2]), int(digits[2:4]), int(digits[4:8])
        else:
            raise ValueError("Data incompleta.")
        return date(year, month, day)

    def _auto_format_date_text(self, line_edit, text):
        digits = "".join(ch for ch in text if ch.isdigit())[:8]

        parts = []
        if digits:
            parts.append(digits[0:2])
        if len(digits) > 2:
            parts.append(digits[2:4])
        if len(digits) > 4:
            parts.append(digits[4:8])

        line_edit.blockSignals(True)
        line_edit.setText("/".join(parts))
        line_edit.blockSignals(False)

    def _finish_date_input(self, line_edit, on_valid=None):
        digits = "".join(ch for ch in line_edit.text() if ch.isdigit())

        try:
            parsed_date = self._parse_date_digits(digits)
        except ValueError:
            QMessageBox.warning(
                self,
                "Data invalida",
                "Digite a data completa: dia, mes e ano (ex: 140826 ou 14/08/2026)."
            )
            line_edit.setText(QDate.currentDate().toString("dd/MM/yyyy"))
            return

        line_edit.blockSignals(True)
        line_edit.setText(parsed_date.strftime("%d/%m/%Y"))
        line_edit.blockSignals(False)

        if on_valid:
            on_valid(parsed_date)

    def _date_value(self, line_edit) -> date:
        digits = "".join(ch for ch in line_edit.text() if ch.isdigit())
        return self._parse_date_digits(digits)

    def add_exam_to_patient(self):
        if not self.current_procedure:
            QMessageBox.warning(self, "Procedimento", "Selecione um procedimento na lista antes de adicionar.")
            return

        try:
            paid_value = Decimal(self.paid_value_input.text().replace(",", "."))
        except InvalidOperation:
            QMessageBox.warning(self, "Valor invalido", "Informe um valor pago valido.")
            return

        if not self.sps_code_input.text().strip():
            QMessageBox.warning(self, "Campos obrigatorios", "Preencha o SAVE do exame.")
            return

        try:
            attendance_date = self._date_value(self.attendance_date_input)
        except ValueError:
            QMessageBox.warning(
                self,
                "Data invalida",
                "Digite a data completa do atendimento (ex: 140826 ou 14/08/2026)."
            )
            return

        self.pending_exams.append(
            {
                "attendance_id": self.editing_pending_exam_id,
                "procedure_id": self.current_procedure.id,
                "procedure_name": self.current_procedure.name,
                "selected_price_type": self.price_type_combo.currentText(),
                "paid_value": paid_value,
                "sps_code": self.sps_code_input.text().strip(),
                "attendance_date": attendance_date,
                "report_month": self.report_month_combo.currentData(),
                "report_year": self.report_year_combo.currentData(),
            }
        )
        self.editing_pending_exam_id = None

        self._refresh_pending_exams_table()

        self.current_procedure = None
        self.procedure_search_input.clear()
        self.procedure_name_label.setText("-")
        self.procedure_source_label.setText("-")
        self.particular_value_label.setText("-")
        self.partner_value_label.setText("-")
        self.premium_value_label.setText("-")
        self.selected_price_label.setText("-")
        self.paid_value_input.clear()
        self.sps_code_input.clear()

    def _refresh_pending_exams_table(self):
        self.exams_table.setRowCount(len(self.pending_exams))
        for row, exam in enumerate(self.pending_exams):
            values = [
                exam["procedure_name"],
                exam["selected_price_type"],
                self.format_money(exam["paid_value"]),
                exam["sps_code"],
                exam["attendance_date"].strftime("%d/%m/%Y"),
            ]
            for column, value in enumerate(values):
                self.exams_table.setItem(row, column, QTableWidgetItem(value))
        self.exams_table.resizeColumnsToContents()

    def remove_selected_exam(self):
        row = self.exams_table.currentRow()
        if row < 0:
            return
        del self.pending_exams[row]
        self._refresh_pending_exams_table()

    def start_edit_attendance(self, attendance):
        self.editing_mode = True
        self.editing_attendance_id = attendance.id
        self.editing_patient_name = attendance.patient_name

        self.patient_input.setText(attendance.patient_name)

        patient_exams = [
            item for item in self.current_rows
            if item.patient_name == attendance.patient_name
        ]
        if attendance.id not in {item.id for item in patient_exams}:
            patient_exams.append(attendance)

        self.original_exam_ids = {item.id for item in patient_exams}

        self.pending_exams = [
            {
                "attendance_id": item.id,
                "procedure_id": item.procedure.id,
                "procedure_name": item.procedure.name,
                "selected_price_type": item.selected_price_type,
                "paid_value": item.paid_value,
                "sps_code": item.sps_code,
                "attendance_date": item.attendance_date,
                "report_month": item.report_month or item.attendance_date.month,
                "report_year": item.report_year or item.attendance_date.year,
            }
            for item in patient_exams
        ]
        self._refresh_pending_exams_table()

        self.editing_pending_exam_id = None
        self.current_procedure = None
        self.procedure_search_input.clear()
        self.procedure_name_label.setText("-")
        self.procedure_source_label.setText("-")
        self.particular_value_label.setText("-")
        self.partner_value_label.setText("-")
        self.premium_value_label.setText("-")
        self.selected_price_label.setText("-")
        self.paid_value_input.clear()
        self.sps_code_input.clear()
        self._reset_attendance_date_to_today()

        self.save_button.setText("Salvar alteracoes")
        self.cancel_edit_button.setVisible(True)
        self.edit_mode_label.setText(
            f"Editando o registro de {attendance.patient_name}. "
            "Adicione, edite (duplo-clique num exame da lista) ou remova exames, "
            "depois clique em \"Salvar alteracoes\"."
        )
        self.edit_mode_label.setVisible(True)

        self.tabs.setCurrentWidget(self.attendance_tab)

    def edit_pending_exam(self, row, column=0):
        if row < 0 or row >= len(self.pending_exams):
            return

        exam = self.pending_exams.pop(row)
        self._refresh_pending_exams_table()

        with SessionLocal() as db:
            procedure = db.get(Procedure, exam["procedure_id"])
            db.expunge(procedure)

        self._display_procedure(procedure)

        price_index = self.price_type_combo.findText(exam["selected_price_type"])
        if price_index >= 0:
            self.price_type_combo.setCurrentIndex(price_index)

        self.attendance_date_input.setText(exam["attendance_date"].strftime("%d/%m/%Y"))

        month_index = self.report_month_combo.findData(exam["report_month"])
        if month_index >= 0:
            self.report_month_combo.setCurrentIndex(month_index)

        year_index = self.report_year_combo.findData(exam["report_year"])
        if year_index == -1:
            self.report_year_combo.addItem(str(exam["report_year"]), exam["report_year"])
            year_index = self.report_year_combo.findData(exam["report_year"])
        self.report_year_combo.setCurrentIndex(year_index)

        self.paid_value_input.setText(str(exam["paid_value"]))
        self.sps_code_input.setText(exam["sps_code"])

        self.editing_pending_exam_id = exam.get("attendance_id")

    def cancel_edit_attendance(self):
        self.editing_mode = False
        self.editing_attendance_id = None
        self.editing_patient_name = None
        self.original_exam_ids = set()
        self.editing_pending_exam_id = None
        self._clear_attendance_form()

        self.add_exam_button.setEnabled(True)
        self.save_button.setText("Salvar atendimentos do paciente")
        self.cancel_edit_button.setVisible(False)
        self.edit_mode_label.setVisible(False)
        self.edit_mode_label.setText("")

    def _reset_attendance_date_to_today(self):
        today = QDate.currentDate()
        self.attendance_date_input.setText(today.toString("dd/MM/yyyy"))
        self.sync_report_month_year(today)

    def _clear_attendance_form(self):
        self.patient_input.clear()
        self.current_procedure = None
        self.procedure_search_input.clear()
        self.procedure_name_label.setText("-")
        self.procedure_source_label.setText("-")
        self.particular_value_label.setText("-")
        self.partner_value_label.setText("-")
        self.premium_value_label.setText("-")
        self.selected_price_label.setText("-")
        self.paid_value_input.clear()
        self.sps_code_input.clear()
        self._reset_attendance_date_to_today()
        self.pending_exams = []
        self._refresh_pending_exams_table()

    def save_patient_attendances(self):
        if self.editing_mode:
            self._save_attendance_edit()
            return

        patient_name = self.patient_input.text().strip()
        if not patient_name:
            QMessageBox.warning(self, "Paciente", "Informe o nome do paciente.")
            return

        if not self.pending_exams:
            QMessageBox.warning(self, "Exames", "Adicione ao menos um exame antes de salvar.")
            return

        try:
            for exam in self.pending_exams:
                create_attendance(
                    patient_name=patient_name,
                    attendance_date=exam["attendance_date"],
                    paid_value=exam["paid_value"],
                    sps_code=exam["sps_code"],
                    procedure_id=exam["procedure_id"],
                    selected_price_type=exam["selected_price_type"],
                    report_month=exam["report_month"],
                    report_year=exam["report_year"],
                )
        except ValueError as exc:
            QMessageBox.warning(self, "Erro", str(exc))
            return

        total_exames = len(self.pending_exams)
        self.pending_exams = []
        self._refresh_pending_exams_table()
        self.patient_input.clear()
        self.refresh_table()
        self.load_report_preview()

        QMessageBox.information(
            self,
            "Atendimentos",
            f"{total_exames} exame(s) de {patient_name} salvos com sucesso."
        )

    def _save_attendance_edit(self):
        patient_name = self.patient_input.text().strip()
        if not patient_name:
            QMessageBox.warning(self, "Paciente", "Informe o nome do paciente.")
            return

        if not self.pending_exams:
            QMessageBox.warning(
                self,
                "Exames",
                "O registro precisa ter ao menos um exame. Adicione um exame ou cancele a edicao."
            )
            return

        current_ids = {
            exam["attendance_id"] for exam in self.pending_exams if exam.get("attendance_id")
        }
        removed_ids = self.original_exam_ids - current_ids

        try:
            for attendance_id in removed_ids:
                delete_attendance(attendance_id)

            for exam in self.pending_exams:
                if exam.get("attendance_id"):
                    update_attendance(
                        attendance_id=exam["attendance_id"],
                        patient_name=patient_name,
                        attendance_date=exam["attendance_date"],
                        paid_value=exam["paid_value"],
                        sps_code=exam["sps_code"],
                        procedure_id=exam["procedure_id"],
                        selected_price_type=exam["selected_price_type"],
                        report_month=exam["report_month"],
                        report_year=exam["report_year"],
                    )
                else:
                    create_attendance(
                        patient_name=patient_name,
                        attendance_date=exam["attendance_date"],
                        paid_value=exam["paid_value"],
                        sps_code=exam["sps_code"],
                        procedure_id=exam["procedure_id"],
                        selected_price_type=exam["selected_price_type"],
                        report_month=exam["report_month"],
                        report_year=exam["report_year"],
                    )
        except ValueError as exc:
            QMessageBox.warning(self, "Erro", str(exc))
            return

        self.editing_mode = False
        self.editing_attendance_id = None
        self.editing_patient_name = None
        self.original_exam_ids = set()
        self.editing_pending_exam_id = None
        self._clear_attendance_form()

        self.save_button.setText("Salvar atendimentos do paciente")
        self.cancel_edit_button.setVisible(False)
        self.edit_mode_label.setVisible(False)
        self.edit_mode_label.setText("")

        self.refresh_table()
        self.load_report_preview()

        QMessageBox.information(self, "Atendimento", "Registro atualizado com sucesso.")

    def refresh_table(self):
        selected_date = None
        if self.search_date_mode.currentIndex() == 1:
            try:
                selected_date = self._date_value(self.search_date_input)
            except ValueError:
                QMessageBox.warning(
                    self,
                    "Data invalida",
                    "Digite a data completa para pesquisar (ex: 140826 ou 14/08/2026)."
                )
                return

        self.current_rows = search_attendances(
            patient=self.search_patient_input.text(),
            procedure=self.search_procedure_input.text(),
            sps_code=self.search_sps_code_input.text(),
            start=selected_date,
            end=selected_date,
        )

        # Agrupa por paciente (mesmo que os ids nao sejam consecutivos, ex: exame novo
        # adicionado durante uma edicao), mantendo o paciente com atividade mais recente
        # no topo, e dentro de cada paciente, o exame mais recente primeiro.
        patient_max_id = {}
        for item in self.current_rows:
            patient_max_id[item.patient_name] = max(patient_max_id.get(item.patient_name, 0), item.id)
        self.current_rows.sort(key=lambda item: (-patient_max_id[item.patient_name], -item.id))

        self.table.setRowCount(len(self.current_rows))
        for row_index, item in enumerate(self.current_rows):
            report_month = item.report_month or item.attendance_date.month
            report_year = item.report_year or item.attendance_date.year
            charged_value = selected_charged_value(item)
            paid_value = item.paid_value
            difference = float(paid_value) - float(charged_value)
            values = [
                item.patient_name,
                item.attendance_date.strftime("%d/%m/%Y"),
                f"{MONTH_NAMES[report_month]} {report_year}",
                item.procedure.code,
                item.procedure.name,
                item.selected_price_type,
                item.procedure.source_sheet,
                self.format_money(charged_value),
                self.format_money(paid_value),
                self.format_money(difference),
                item.sps_code,
            ]
            for column, value in enumerate(values):
                cell = QTableWidgetItem(value)
                if column in (7, 8, 9):
                    cell.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                elif column == 0:
                    cell.setTextAlignment(Qt.AlignCenter)
                self.table.setItem(row_index, column, cell)

        self.table.clearSpans()
        row_index = 0
        total = len(self.current_rows)
        while row_index < total:
            span = 1
            while (
                row_index + span < total
                and self.current_rows[row_index + span].patient_name
                == self.current_rows[row_index].patient_name
            ):
                span += 1
            if span > 1:
                self.table.setSpan(row_index, 0, span, 1)
                self.table.setSpan(row_index, 11, span, 1)

            edit_button = QPushButton("Editar")
            edit_button.setObjectName("tableActionButton")
            edit_button.clicked.connect(
                lambda checked=False, attendance=self.current_rows[row_index]: self.start_edit_attendance(attendance)
            )

            button_container = QWidget()
            button_layout = QHBoxLayout(button_container)
            button_layout.setContentsMargins(10, 6, 10, 6)
            button_layout.addWidget(edit_button)

            self.table.setCellWidget(row_index, 11, button_container)

            row_index += span

        self.table.setVerticalHeaderLabels(
            [str(total - row_index) for row_index in range(total)]
        )

        self.table.resizeColumnsToContents()

    def export_current_excel(self):
        resolved = self._resolve_report_range()
        if not resolved:
            return
        mes_inicial, ano_inicial, mes_final, ano_final, rows = resolved

        path = export_excel(
            rows,
            filename=self._report_range_filename(mes_inicial, ano_inicial, mes_final, ano_final, "xlsx")
        )

        QMessageBox.information(
            self,
            "Excel gerado",
            f"Arquivo salvo em:\n{path}"
        )
    def export_current_pdf(self):
        if self.current_preview_kind == "budget":
            if not self.budget_items:
                QMessageBox.information(
                    self,
                    "Orçamento",
                    "Nao ha orcamento para baixar. Adicione procedimentos na aba Orcamento."
                )
                return

            path = export_budget_pdf(self.budget_items, filename="Orcamento.pdf")

            QMessageBox.information(
                self,
                "PDF gerado",
                f"Arquivo salvo em:\n{path}"
            )
            return

        resolved = self._resolve_report_range()
        if not resolved:
            return
        mes_inicial, ano_inicial, mes_final, ano_final, rows = resolved

        path = export_pdf(
            rows,
            filename=self._report_range_filename(mes_inicial, ano_inicial, mes_final, ano_final, "pdf")
        )

        QMessageBox.information(
            self,
            "PDF gerado",
            f"Arquivo salvo em:\n{path}"
        )

    def print_current_report(self):
        temp_pdf = Path(tempfile.gettempdir()) / "cpp_print.pdf"

        if self.current_preview_kind == "budget":
            if not self.budget_items:
                QMessageBox.information(
                    self,
                    "Orçamento",
                    "Nao ha orcamento para imprimir. Adicione procedimentos na aba Orcamento."
                )
                return
            export_budget_pdf(self.budget_items, filename=str(temp_pdf))
        else:
            resolved = self._resolve_report_range()
            if not resolved:
                return
            _, _, _, _, rows = resolved
            export_pdf(rows, filename=str(temp_pdf))

        document = QPdfDocument(self)
        document.load(str(temp_pdf))

        if document.pageCount() == 0:
            QMessageBox.warning(
                self,
                "Imprimir",
                "Nao foi possivel preparar o relatorio para impressao."
            )
            return

        preview = PrintPreviewDialog(document, self)
        preview.exec()

    def search_procedures(self):
        texto = self.procedure_search_input.text().strip()

        self.procedure_list.clear()

        if len(texto) < 2:
            return

        with SessionLocal() as db:
            resultados = (
                db.query(Procedure)
                .filter(
                    or_(
                        Procedure.name.ilike(f"%{texto}%"),
                        Procedure.code.ilike(f"%{texto}%")
                    )
                )
                .order_by(Procedure.name, Procedure.source_sheet)
                .all()
            )

        # Conta os prestadores/médicos distintos de cada código na própria
        # planilha. Só exibimos o nome quando o código realmente é compartilhado.
        providers_by_code = self.price_table_providers_by_code

        def providers_for(proc):
            code = "" if proc.code is None else str(proc.code).strip()
            digits = "".join(ch for ch in code if ch.isdigit())
            source = (getattr(proc, "source_sheet", "") or "").strip()
            description_key = _normalize_lookup_text(proc.name)

            def money_key(value):
                if value in (None, ""):
                    return None
                try:
                    return round(float(value), 6)
                except (TypeError, ValueError):
                    return None

            price_signature = (
                money_key(getattr(proc, "reference_value", None)),
                money_key(getattr(proc, "partner_value", None)),
                money_key(getattr(proc, "premium_value", None)),
            )

            exact = self.price_table_provider_by_key.get(
                (source, digits, description_key, price_signature), set()
            )
            if exact:
                return sorted(exact)

            source_providers = self.price_table_provider_by_source_code.get(
                (source, digits), set()
            )
            if source_providers:
                return sorted(source_providers)

            return [source] if source else []

        for proc in resultados:
            codigo = proc.code if proc.code else "SEM CÓDIGO"
            codigo_normalizado = "".join(
                ch for ch in str(proc.code or "") if ch.isdigit()
            )
            prestadores = providers_by_code.get(codigo_normalizado, set())

            # Código único: uma única linha, exatamente no padrão original.
            if not codigo_normalizado or len(prestadores) <= 1:
                display_items = [(proc, None)]
            else:
                # IMPORTANTE:
                # um mesmo registro do banco pode representar vários blocos
                # da planilha quando código + descrição + preços são idênticos.
                # Nesse caso NÃO eliminamos os prestadores: criamos uma linha
                # visual para cada um, reutilizando o mesmo procedimento do banco.
                provider_list = providers_for(proc)
                if not provider_list:
                    provider_list = [None]
                display_items = [(proc, provider) for provider in provider_list]

            for display_proc, provider in display_items:
                item_text = f"{codigo} - {display_proc.name}"
                if provider:
                    item_text += f" - {provider}"

                item = QListWidgetItem(item_text)
                item.setData(Qt.ItemDataRole.UserRole, display_proc)
                self.procedure_list.addItem(item)

    def select_procedure(self, item):
        procedure = item.data(Qt.UserRole)
        self._display_procedure(procedure)

def run() -> int:
    bootstrap_database()
    app = QApplication(sys.argv)
    app.setStyleSheet(APP_STYLESHEET)

    while True:
        login = LoginDialog()
        if login.exec() != QDialog.Accepted:
            return 0

        window = MainWindow(login.user)
        window.show()
        app.exec()

        if not window.logged_out:
            return 0
        # senao, volta para o topo do laco e mostra a tela de login de novo