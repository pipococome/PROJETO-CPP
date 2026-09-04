from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


def _normalize(value) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).upper()


def _code(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip().replace(".0", "")
    digits = re.sub(r"\D", "", text)
    if len(digits) >= 6:
        return digits
    return None


def _money(value) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _looks_like_header(value) -> bool:
    text = _normalize(value)
    return any(
        word in text
        for word in [
            "PROCEDIMENTO",
            "CODIGO",
            "TUSS",
            "DESCRICAO",
            "PARTICULAR",
            "PARCEIRO",
            "PREMIUM",
            "CREDENCIADO",
        ]
    )


def _find_price_columns(row, description_index: int, header_labels: dict[int, str]) -> dict[str, Decimal]:
    prices = {}
    fallback = None
    for index in range(description_index + 1, len(row)):
        amount = _money(row[index])
        if amount is None:
            continue
        header_text = header_labels.get(index, "")
        if "PARTICULAR" in header_text:
            prices.setdefault("reference_value", amount)
        elif "PARCEIRO" in header_text:
            prices.setdefault("partner_value", amount)
        elif "PREMIUM" in header_text or "PRO" in header_text:
            prices.setdefault("premium_value", amount)
        elif fallback is None:
            fallback = amount
    if "reference_value" not in prices and fallback is not None:
        prices["reference_value"] = fallback
    return prices


def extract_procedures(path: str | Path) -> list[dict]:
    workbook_path = Path(path)

    if not workbook_path.exists():
        return []

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    procedures = []
    seen = set()

    for sheet in workbook.worksheets:

        header_labels = {}

        for row in sheet.iter_rows(values_only=True):

            if not any(cell is not None for cell in row):
                continue

            # Atualiza cabeçalhos
            for col, cell in enumerate(row):
                text = _normalize(cell)
                if _looks_like_header(text):
                    header_labels[col] = text

            code = None
            description = None
            description_index = None

            # =====================================================
            # Procura um código na linha
            # =====================================================

            for col, cell in enumerate(row):

                found = _code(cell)

                if found:

                    code = found

                    for j in range(col + 1, min(col + 5, len(row))):

                        value = row[j]

                        if (
                            isinstance(value, str)
                            and value.strip()
                            and not _looks_like_header(value)
                        ):
                            description = value.strip()
                            description_index = j
                            break

                    break

            # =====================================================
            # Caso NÃO exista código
            # =====================================================

            if description is None:

                for col, cell in enumerate(row):

                    if not isinstance(cell, str):
                        continue

                    text = cell.strip()

                    if (
                        len(text) < 5
                        or _looks_like_header(text)
                    ):
                        continue

                    prices = _find_price_columns(
                        row,
                        col,
                        header_labels,
                    )

                    if "reference_value" not in prices:
                        continue

                    description = text
                    description_index = col
                    code = ""
                    break

            if description is None:
                continue

            prices = _find_price_columns(
                row,
                description_index,
                header_labels,
            )

            if "reference_value" not in prices:
                continue

            key = (
                code,
                description,
                str(prices["reference_value"]),
                sheet.title,
            )

            if key in seen:
                continue

            seen.add(key)

            procedures.append(
                {
                    "code": code,
                    "name": description,
                    "reference_value": prices["reference_value"],
                    "partner_value": prices.get("partner_value"),
                    "premium_value": prices.get("premium_value"),
                    "source_sheet": sheet.title,
                    "price_type": "PARTICULAR",
                }
            )

    workbook.close()

    return procedures