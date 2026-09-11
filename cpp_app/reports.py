from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from cpp_app.config import EXPORT_DIR, REPORT_TEMPLATE_PATH


MONTH_NAMES = {
    1: "JANEIRO",
    2: "FEVEREIRO",
    3: "MAR\u00c7O",
    4: "ABRIL",
    5: "MAIO",
    6: "JUNHO",
    7: "JULHO",
    8: "AGOSTO",
    9: "SETEMBRO",
    10: "OUTUBRO",
    11: "NOVEMBRO",
    12: "DEZEMBRO",
}


REPORT_HEADERS = [
    "Paciente",
    "Data",
    "SAVE",
    "Descricao",
    "Tabela Origem",
    "Valor Cobrado",
    "Valor Pago",
    "Diferenca",
]

def selected_charged_value(item) -> float:
    if getattr(item, "charged_value", None) is not None:
        return float(item.charged_value)
    return float(item.procedure.reference_value)


def format_money(value) -> str:
    if value is None:
        return ""

    value = float(value)

    text = f"{abs(value):,.2f}"
    text = text.replace(",", "X").replace(".", ",").replace("X", ".")

    if value < 0:
        return f"-R$ {text}"

    return f"R$ {text}"


def truncate_text(text: str, max_width: float, font_name: str = "Helvetica", font_size: float = 8) -> str:
    """Corta o texto e adiciona '...' se ele nao couber em max_width (em pontos)."""
    if stringWidth(text, font_name, font_size) <= max_width:
        return text

    ellipsis = "..."
    ellipsis_width = stringWidth(ellipsis, font_name, font_size)

    truncated = text
    while truncated and stringWidth(truncated, font_name, font_size) + ellipsis_width > max_width:
        truncated = truncated[:-1]

    return truncated.rstrip() + ellipsis


def attendances_to_rows(attendances) -> list[list]:
    rows = []

    for item in attendances:
        charged_value = selected_charged_value(item)
        paid_value = float(item.paid_value)
        difference = paid_value - charged_value

        rows.append(
            [
                item.patient_name,
                item.attendance_date.strftime("%d/%m/%Y"),
                item.sps_code,
                item.procedure.name,
                item.procedure.source_sheet,
                format_money(charged_value),
                format_money(paid_value),
                format_money(difference),
            ]
        )

    return rows


def sum_differences(attendances) -> tuple[float, float]:
    """Retorna (total_positivo, total_negativo) das diferencas (Valor Pago - Valor Cobrado)."""
    total_positive = 0.0
    total_negative = 0.0

    for item in attendances:
        charged_value = selected_charged_value(item)
        paid_value = float(item.paid_value)
        difference = paid_value - charged_value

        if difference >= 0:
            total_positive += difference
        else:
            total_negative += difference

    return total_positive, total_negative


def _report_month_year(item) -> tuple[int, int]:
    month = getattr(item, "report_month", None) or item.attendance_date.month
    year = getattr(item, "report_year", None) or item.attendance_date.year
    return month, year


def _month_sheet_name(workbook, month: int) -> str | None:
    """O modelo real nao tem o ano no nome da aba (ex: 'MAIO', 'JANEIRO ' com espaco extra),
    entao a busca e so pelo nome do mes."""
    expected = MONTH_NAMES[month]
    for name in workbook.sheetnames:
        if name.strip().upper() == expected:
            return name
    return None


def _last_data_row(sheet) -> int:
    last = 4
    for row in range(4, sheet.max_row + 1):
        if any(sheet.cell(row=row, column=col).value not in (None, "") for col in range(1, 7)):
            last = row
    return last


def _update_month_totals(sheet) -> None:
    last = max(_last_data_row(sheet), 4)
    sheet["I4"] = f'=SUMIF(G4:G{last},">=0")'
    sheet["I7"] = f'=SUMIF(G4:G{last},"<=0")'
    sheet["I10"] = "=I4+I7"
    sheet["I9"] = '=IF(I10>0,"LUCRO",IF(I10<0,"PREJU\u00cdZO","ZERO"))'


def _rebuild_annual_formulas(workbook, kept_titles: list[str]) -> None:
    """Reescreve as formulas de soma da TOTAL ANUAL pra apontar pros nomes de aba
    finais (com o ano), ja que o modelo original referenciava os nomes sem ano."""
    if "TOTAL ANUAL" not in workbook.sheetnames:
        return

    sheet = workbook["TOTAL ANUAL"]

    if not kept_titles:
        sheet["A5"] = 0
        sheet["E5"] = 0
        return

    positive_refs = ",".join(f"'{title}'!I4:K4" for title in kept_titles)
    negative_refs = ",".join(f"'{title}'!I7:K7" for title in kept_titles)

    sheet["A5"] = f"=SUM({positive_refs})"
    sheet["E5"] = f"=SUM({negative_refs})"


def export_excel(attendances, filename: str = "relatorio_atendimentos.xlsx") -> Path:
    """Gera um arquivo com uma aba por mes/ano presente nos atendimentos, renomeando
    cada aba pra incluir o ano (ex: 'MAIO 2026'), clonando a aba original do modelo
    quando o mesmo mes aparece em mais de um ano na mesma exportacao. A aba TOTAL ANUAL
    e mantida, com as formulas reescritas pra somar as abas com os nomes finais."""
    attendances = list(attendances)
    path = EXPORT_DIR / filename
    template_path = Path(REPORT_TEMPLATE_PATH)

    if not attendances:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Atendimentos"
        worksheet.append(REPORT_HEADERS)
        workbook.save(path)
        return path

    if not template_path.exists():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Atendimentos"
        worksheet.append(REPORT_HEADERS)
        for row in attendances_to_rows(attendances):
            worksheet.append(row)
        for column in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in column)
            worksheet.column_dimensions[column[0].column_letter].width = max_len + 2
        workbook.save(path)
        return path

    # Agrupa os atendimentos por (mes, ano)
    groups: dict[tuple[int, int], list] = {}
    for item in attendances:
        groups.setdefault(_report_month_year(item), []).append(item)

    # Agrupa por mes -> quais anos aparecem, pra saber quando precisa clonar
    months_years: dict[int, list[int]] = {}
    for month, year in groups:
        months_years.setdefault(month, []).append(year)

    workbook = load_workbook(template_path)
    unformatted_items = []
    sheet_for: dict[tuple[int, int], object] = {}

    for month, years in months_years.items():
        base_name = _month_sheet_name(workbook, month)
        if not base_name:
            for year in years:
                unformatted_items.extend(groups[(month, year)])
            continue

        original_sheet = workbook[base_name]

        for index, year in enumerate(sorted(years)):
            final_title = f"{MONTH_NAMES[month]} {year}"[:31]

            if index == 0:
                # primeira vez que esse mes aparece: reaproveita a aba original do modelo
                sheet = original_sheet
                sheet.title = final_title
            else:
                # mesmo mes, outro ano: clona a aba original (ainda vazia nesse ponto)
                sheet = workbook.copy_worksheet(original_sheet)
                sheet.title = final_title

            sheet_for[(month, year)] = sheet

    kept = []  # lista de (ano, mes, titulo) pra depois ordenar cronologicamente

    for (month, year), items in groups.items():
        sheet = sheet_for.get((month, year))
        if not sheet:
            continue

        items_sorted = sorted(items, key=lambda item: item.attendance_date)

        for offset, item in enumerate(items_sorted):
            row = 4 + offset
            sheet.cell(row=row, column=1).value = item.attendance_date
            sheet.cell(row=row, column=2).value = item.sps_code
            sheet.cell(row=row, column=3).value = item.patient_name
            sheet.cell(row=row, column=4).value = item.procedure.name
            sheet.cell(row=row, column=5).value = selected_charged_value(item)
            sheet.cell(row=row, column=6).value = float(item.paid_value)
            sheet.cell(row=row, column=7).value = f"=F{row}-E{row}"

            sheet.cell(row=row, column=1).number_format = "dd/mm/yyyy"
            for col in (5, 6, 7):
                sheet.cell(row=row, column=col).number_format = 'R$ #,##0.00'

        _update_month_totals(sheet)
        kept.append((year, month, sheet.title))

    kept.sort()
    kept_sheet_titles = [title for _, _, title in kept]

    # Mantem no arquivo somente as abas usadas (mais o TOTAL ANUAL), removendo as demais
    for name in list(workbook.sheetnames):
        if name not in kept_sheet_titles and name != "TOTAL ANUAL":
            del workbook[name]

    # Reordena cronologicamente, deixando o TOTAL ANUAL sempre por ultimo
    ordered_titles = list(kept_sheet_titles)
    if "TOTAL ANUAL" in workbook.sheetnames:
        ordered_titles.append("TOTAL ANUAL")
    workbook._sheets = [workbook[title] for title in ordered_titles]

    _rebuild_annual_formulas(workbook, kept_sheet_titles)

    if unformatted_items:
        # Meses sem aba correspondente no modelo (nao deveria acontecer com um modelo
        # de 12 meses normal, mas fica como rede de seguranca)
        cover = workbook.create_sheet("SEM FORMATO")
        cover.append(REPORT_HEADERS)
        for row in attendances_to_rows(unformatted_items):
            cover.append(row)
        for column in cover.columns:
            max_len = max(len(str(cell.value or "")) for cell in column)
            cover.column_dimensions[column[0].column_letter].width = max_len + 2

    workbook.save(path)
    return path


def export_pdf(attendances, filename: str = "relatorio_atendimentos.pdf") -> Path:
    path = EXPORT_DIR / filename

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    story = [
        Paragraph(
            "PLANILHA VALORES NÃO COBRADO - PA/ELETIVO",
            styles["Title"]
        ),
        Spacer(1, 12),
    ]

    attendances = list(attendances)

    data = [REPORT_HEADERS]

    for row in attendances_to_rows(attendances):
        data.append([str(value) for value in row])

    col_widths = [
        120,  # Paciente
        60,   # Data
        60,   # SAVE
        210,  # Descrição
        90,   # Tabela Origem
        80,   # Valor Cobrado
        80,   # Valor Pago
        80,   # Diferença
    ]

    cell_padding = 12  # LEFTPADDING + RIGHTPADDING padrao do reportlab

    for row in data[1:]:
        for column in (0, 2, 3, 4):  # Paciente, SAVE, Descricao, Tabela Origem
            max_width = col_widths[column] - cell_padding
            row[column] = truncate_text(row[column], max_width)

    table = Table(
        data,
        colWidths=col_widths,
        repeatRows=1,
    )
    table.hAlign = "LEFT"

    table_style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        # Nome do paciente sempre centralizado, tenha ele 1 exame ou varios.
        ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ("VALIGN", (0, 1), (0, -1), "MIDDLE"),
    ]

    row_index = 1  # linha 0 e o cabecalho
    i = 0
    total_rows = len(attendances)
    while i < total_rows:
        j = i
        while j + 1 < total_rows and attendances[j + 1].patient_name == attendances[i].patient_name:
            j += 1
        span = j - i
        if span > 0:
            table_style_commands.append(
                ("SPAN", (0, row_index), (0, row_index + span))
            )
        row_index += span + 1
        i = j + 1

    table.setStyle(TableStyle(table_style_commands))

    story.append(table)

    total_positive, total_negative = sum_differences(attendances)

    story.append(Spacer(1, 18))

    summary_data = [
        ["Total Recebido a Mais", "Total Nao Cobrado"],
        [format_money(total_positive), format_money(total_negative)],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[160, 160],
    )
    summary_table.hAlign = "LEFT"

    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 12),
                ("TEXTCOLOR", (0, 1), (0, 1), colors.HexColor("#1a7a1a")),
                ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#b30000")),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )

    story.append(summary_table)

    net_result = total_positive + total_negative

    if net_result > 0:
        result_title = "LUCRO"
        result_color = colors.HexColor("#1a7a1a")
    elif net_result < 0:
        result_title = "PREJUIZO"
        result_color = colors.HexColor("#b30000")
    else:
        result_title = "ZERO"
        result_color = colors.black

    story.append(Spacer(1, 18))

    result_data = [
        [result_title],
        [format_money(net_result)],
    ]

    result_table = Table(
        result_data,
        colWidths=[320],
    )
    result_table.hAlign = "LEFT"

    result_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 1), (-1, 1), 12),
                ("TEXTCOLOR", (0, 1), (0, 1), result_color),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )

    story.append(result_table)

    doc.build(story)

    return path


BUDGET_HEADERS = ["Codigo", "Procedimento", "Particular", "Parceiros", "Pro Premium"]


def export_budget_pdf(procedures, filename: str = "orcamento.pdf") -> Path:
    """Gera um PDF de orcamento com os procedimentos selecionados, mostrando os 3 tipos
    de valor (Particular/Parceiros/Pro Premium) e a soma total de cada um no final."""
    procedures = list(procedures)
    path = EXPORT_DIR / filename

    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    story = [Paragraph("OR\u00c7AMENTO", styles["Title"]), Spacer(1, 12)]

    col_widths = [90, 380, 100, 100, 100]

    data = [BUDGET_HEADERS]

    total_particular = 0.0
    total_partner = 0.0
    total_premium = 0.0

    for procedure in procedures:
        codigo = procedure.code if procedure.code else "SEM C\u00d3DIGO"
        codigo = truncate_text(codigo, col_widths[0] - 12, font_size=9)
        nome = truncate_text(procedure.name, col_widths[1] - 12, font_size=9)

        data.append(
            [
                codigo,
                nome,
                format_money(procedure.reference_value),
                format_money(procedure.partner_value),
                format_money(procedure.premium_value),
            ]
        )

        total_particular += float(procedure.reference_value or 0)
        total_partner += float(procedure.partner_value or 0)
        total_premium += float(procedure.premium_value or 0)

    data.append(
        [
            "",
            "TOTAL",
            format_money(total_particular),
            format_money(total_partner),
            format_money(total_premium),
        ]
    )

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.hAlign = "LEFT"

    last_row = len(data) - 1

    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f4e79")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 1), (-1, last_row - 1), colors.white),
                ("BACKGROUND", (0, last_row), (-1, last_row), colors.HexColor("#eef3fa")),
                ("FONTNAME", (0, last_row), (-1, last_row), "Helvetica-Bold"),
                ("SPAN", (0, last_row), (1, last_row)),
                ("ALIGN", (0, last_row), (1, last_row), "RIGHT"),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ]
        )
    )

    story.append(table)

    doc.build(story)

    return path